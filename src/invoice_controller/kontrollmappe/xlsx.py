"""Kontrollmappe .xlsx writer (R6) — the PLAN.md sheet set, lean v1.

Sheets: Übersicht, Angebote, Rechnungen, Abgleich (the core), Prüfungen
(Datum/Adresse), Summen, Audit-Log. Q1 (2026-09-03): in the Abgleich any variance
that is not exactly zero renders RED — no tolerance bands, the consultant decides.
Q2: no checkbox column. Live SUM formulas where sums span rows; review flags share
the color language of the other writers (red = look, orange/yellow = note).
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from invoice_controller.kontrollmappe.build import KontrollmappeResult
from invoice_controller.match.vendor import normalize_vendor
from invoice_controller.models import InvoiceDocument

_EUR = "#,##0.00"
_PCT = "0.0%"
_DATE = "DD.MM.YYYY"
_RED = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
_YELLOW = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_GREEN = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _eur(ws: Worksheet, r: int, c: int, v: Decimal | None):
    cell = ws.cell(r, c)
    if v is not None:
        cell.value = float(v)
        cell.number_format = _EUR
    return cell


def _headers(ws: Worksheet, r: int, labels: list[str]) -> None:
    for c, text in enumerate(labels, start=1):
        ws.cell(r, c, text).font = _BOLD


def _widths(ws: Worksheet, spec: dict[str, int]) -> None:
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def _uebersicht(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.active
    ws.title = "Übersicht"
    ws.cell(1, 1, "Kontrollmappe — Übersicht").font = _BOLD
    rows = [
        ("Kunde", result.config.client.name if result.config.client else None),
        ("Bewilligungszeitraum",
         f"{result.config.bewilligungszeitraum_start:%d.%m.%Y} – {result.config.bewilligungszeitraum_end:%d.%m.%Y}"
         if result.config.has_window else None),
        ("Lieferanten (Angebote)", len(result.groups)),
        ("Rechnungen zugeordnet", sum(len(g.invoices) for g in result.groups)),
        ("Rechnungen ohne Angebots-Lieferant", len(result.offerless_invoices)),
        ("nicht verarbeitbare Dateien", len(result.unreadable)),
    ]
    r = 3
    for label, value in rows:
        ws.cell(r, 1, label)
        cell = ws.cell(r, 2, value if value is not None else "fehlt")
        if value is None:
            cell.fill = _RED
        r += 1
    r += 1
    for flag in result.flags:
        ws.cell(r, 1, f"⚠ {flag}").fill = _YELLOW
        r += 1
    for path, reason in result.unreadable:
        ws.cell(r, 1, f"✗ {path.name}: {reason}").fill = _RED
        r += 1
    _widths(ws, {"A": 42, "B": 46})


def _angebote(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Angebote")
    _headers(ws, 1, ["Lieferant", "Angebot", "Datum", "Pos", "Beschreibung", "Netto €", "Optional", "Kategorie"])
    r = 2
    for g in result.groups:
        for offer in g.offers:
            for pos in offer.positions:
                ws.cell(r, 1, offer.header.vendor_name)
                ws.cell(r, 2, offer.header.offer_number)
                ws.cell(r, 3, offer.header.offer_date).number_format = _DATE
                ws.cell(r, 4, pos.pos)
                ws.cell(r, 5, pos.description).alignment = _WRAP
                _eur(ws, r, 6, pos.line_total_net)
                ws.cell(r, 7, "optional" if pos.optional else None)
                ws.cell(r, 8, pos.kategorie.value)
                r += 1
    _widths(ws, {"A": 26, "B": 14, "C": 11, "D": 7, "E": 52, "F": 12, "G": 9, "H": 16})


def _rechnungen(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Rechnungen")
    _headers(ws, 1, ["Lieferant", "Re-Nr.", "Datum", "Typ", "Netto €", "Brutto €",
                     "Kreuzsumme", "Prüfhinweise"])
    r = 2
    all_invoices = [i for g in result.groups for i in g.invoices] + result.offerless_invoices
    for inv in all_invoices:
        ws.cell(r, 1, inv.vendor_name)
        ws.cell(r, 2, inv.invoice_number)
        ws.cell(r, 3, inv.invoice_date).number_format = _DATE
        ws.cell(r, 4, inv.invoice_type.value)
        _eur(ws, r, 5, inv.netto)
        _eur(ws, r, 6, inv.brutto)
        cs = inv.position_check
        cell = ws.cell(r, 7, "OK" if cs and cs.passed else "⚠")
        if not (cs and cs.passed):
            cell.fill = _RED
        problems = []
        if cs and not cs.passed and cs.message:
            problems.append(cs.message)
        if not inv.amount_check.passed and inv.amount_check.message:
            problems.append(inv.amount_check.message)
        if inv.grounding_check is not None and not inv.grounding_check.passed:
            problems.append(inv.grounding_check.message)
        note = ws.cell(r, 8, "; ".join(p for p in problems if p) or None)
        note.alignment = _WRAP
        if problems:
            note.fill = _YELLOW
        r += 1
    _widths(ws, {"A": 26, "B": 16, "C": 11, "D": 14, "E": 12, "F": 12, "G": 11, "H": 52})


def _abgleich(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Abgleich")
    _headers(ws, 1, ["Pos", "Beschreibung (Angebot)", "angeboten €", "abgerechnet €",
                     "Abweichung €", "Abw. %", "Konfidenz", "Rechnungszeilen", "Hinweis"])
    r = 3
    for g in result.groups:
        ws.cell(r, 1, f"Lieferant: {g.label}").font = _BOLD
        r += 1
        for row in g.rows:
            ws.cell(r, 1, row.offer_position.pos)
            desc = row.offer_position.description
            if row.offer_position.optional:
                desc += " (optional)"
            ws.cell(r, 2, desc).alignment = _WRAP
            _eur(ws, r, 3, row.offered)
            _eur(ws, r, 4, row.invoiced if row.matched else None)
            refs = "; ".join(
                f"{m.invoice.invoice_number}#{m.invoice.positions[m.position_index].pos or m.position_index + 1}"
                for m in row.matched
            )
            notes = "; ".join(dict.fromkeys(m.note for m in row.matched if m.note))
            conf = min((m.confidence.value for m in row.matched), default=None)
            if not row.matched:
                if not row.offer_position.optional:
                    ws.cell(r, 9, "FEHLT — keine Rechnungszeile zugeordnet").fill = _RED
                else:
                    ws.cell(r, 9, "optional, nicht abgerufen")
            else:
                variance = row.variance
                _eur(ws, r, 5, variance)
                if row.offered not in (None, Decimal(0)):
                    pct = ws.cell(r, 6, float(variance / row.offered))
                    pct.number_format = _PCT
                # Q1 (2026-09-03): ANY variance ≠ 0 → red, consultant decides.
                if variance is not None and variance != 0:
                    ws.cell(r, 5).fill = _RED
                    ws.cell(r, 6).fill = _RED
                ws.cell(r, 7, conf)
                if conf == "medium":
                    ws.cell(r, 7).fill = _YELLOW
                ws.cell(r, 8, refs).alignment = _WRAP
                if notes:
                    ws.cell(r, 9, notes).alignment = _WRAP
            r += 1
        for extra in g.extras:
            inv_pos = extra.invoice.positions[extra.position_index]
            ws.cell(r, 1, "EXTRA").fill = _RED
            ws.cell(r, 2, inv_pos.description).alignment = _WRAP
            _eur(ws, r, 4, extra.amount)
            ws.cell(r, 8, f"{extra.invoice.invoice_number}#{inv_pos.pos or extra.position_index + 1}")
            ws.cell(r, 9, extra.note or "keine Angebotsposition zugeordnet").fill = _RED
            r += 1
        r += 1
    if result.offerless_invoices:
        ws.cell(r, 1, "Rechnungen ohne Angebots-Lieferant").font = _BOLD
        r += 1
        for inv in result.offerless_invoices:
            ws.cell(r, 2, f"{inv.vendor_name} — {inv.invoice_number}")
            _eur(ws, r, 4, inv.netto)
            ws.cell(r, 9, "extra — kein Angebot dieses Lieferanten").fill = _RED
            r += 1
    _widths(ws, {"A": 10, "B": 46, "C": 12, "D": 13, "E": 12, "F": 9, "G": 10, "H": 26, "I": 40})


def _address_ok(inv: InvoiceDocument, client_name: str) -> bool | None:
    if inv.recipient_name is None:
        return None
    client = normalize_vendor(client_name)
    recipient = normalize_vendor(inv.recipient_name)
    return len(client & recipient) > 0


def _pruefungen(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Prüfungen")
    _headers(ws, 1, ["Lieferant", "Re-Nr.", "Re-Datum", "im Bewilligungszeitraum",
                     "Empfänger lt. Rechnung", "Empfänger = Kunde"])
    config = result.config
    r = 2
    for inv in [i for g in result.groups for i in g.invoices] + result.offerless_invoices:
        ws.cell(r, 1, inv.vendor_name)
        ws.cell(r, 2, inv.invoice_number)
        ws.cell(r, 3, inv.invoice_date).number_format = _DATE
        window = config.window_ok(inv.invoice_date)
        cell = ws.cell(r, 4, {True: "ja", False: "NEIN", None: ""}[window])
        if window is False:
            cell.fill = _RED
        ws.cell(r, 5, inv.recipient_name if not inv.masked else "[maskiert]").alignment = _WRAP
        if config.client is not None:
            # A1: masked runs carry the deterministic pre-masking check.
            ok = inv.recipient_local_ok if inv.recipient_local_ok is not None else _address_ok(inv, config.client.name)
            acell = ws.cell(r, 6, {True: "ja", False: "NEIN", None: "nicht lesbar"}[ok])
            if ok is False:
                acell.fill = _RED
            elif ok is None:
                acell.fill = _YELLOW
        r += 1
    _widths(ws, {"A": 26, "B": 16, "C": 11, "D": 20, "E": 34, "F": 16})


def _summen(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Summen")
    _headers(ws, 1, ["Lieferant", "angeboten €", "abgerechnet €", "Abweichung €"])
    r = 2
    first, last = r, r
    for g in result.groups:
        ws.cell(r, 1, g.label)
        _eur(ws, r, 2, g.offered_total)
        _eur(ws, r, 3, g.invoiced_total)
        diff = g.invoiced_total - g.offered_total
        cell = _eur(ws, r, 4, diff)
        if diff != 0:
            cell.fill = _RED
        last = r
        r += 1
    for inv in result.offerless_invoices:
        ws.cell(r, 1, f"{inv.vendor_name} (ohne Angebot)").fill = _YELLOW
        _eur(ws, r, 3, inv.netto)
        last = r
        r += 1
    ws.cell(r + 1, 1, "Gesamt").font = _BOLD
    for col_letter, col in (("B", 2), ("C", 3), ("D", 4)):
        ws.cell(r + 1, col, f"=SUM({col_letter}{first}:{col_letter}{last})").number_format = _EUR
    _widths(ws, {"A": 34, "B": 14, "C": 14, "D": 14})


def _audit(wb: Workbook, result: KontrollmappeResult) -> None:
    ws = wb.create_sheet("Audit-Log")
    _headers(ws, 1, ["Dokument", "Extraktionsweg", "Kreuzsumme", "Betragsprüfung", "Verankerung"])
    r = 2
    docs = [(o.source_path.name, o.extraction_method, o.cross_sum.passed, None,
             o.grounding_check.passed if o.grounding_check else None)
            for g in result.groups for o in g.offers]
    docs += [(i.source_path.name, i.extraction_method,
              i.position_check.passed if i.position_check else None,
              i.amount_check.passed,
              i.grounding_check.passed if i.grounding_check else None)
             for g in result.groups for i in g.invoices]
    docs += [(i.source_path.name, i.extraction_method,
              i.position_check.passed if i.position_check else None,
              i.amount_check.passed,
              i.grounding_check.passed if i.grounding_check else None)
             for i in result.offerless_invoices]
    for name, method, cs, ac, gc in docs:
        ws.cell(r, 1, name)
        cell = ws.cell(r, 2, method)
        if "+retry" in method:
            cell.fill = _YELLOW
        for col, val in ((3, cs), (4, ac), (5, gc)):
            c = ws.cell(r, col, {True: "OK", False: "⚠", None: "—"}[val])
            if val is False:
                c.fill = _RED
        r += 1
    _widths(ws, {"A": 52, "B": 22, "C": 12, "D": 14, "E": 12})


def write_kontrollmappe(result: KontrollmappeResult, output_path: Path) -> None:
    wb = Workbook()
    _uebersicht(wb, result)
    _angebote(wb, result)
    _rechnungen(wb, result)
    _abgleich(wb, result)
    _pruefungen(wb, result)
    _summen(wb, result)
    _audit(wb, result)
    wb.save(output_path)


def kontrollmappe_filename(projekt: str, on: date | None = None) -> str:
    """Q3 (2026-09-03): Kontrollmappe_<Projekt>_<YYYY-MM-DD>.xlsx."""
    return f"Kontrollmappe_{projekt.strip().replace('/', '-')}_{(on or date.today()):%Y-%m-%d}.xlsx"
