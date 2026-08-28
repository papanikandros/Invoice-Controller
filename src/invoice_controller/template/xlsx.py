"""Kostenaufstellung writer — .xlsx (openpyxl), the Microsoft-native successor of
template/ods.py (format decision 2026-08-28: end users are on Windows).

Block shape is identical to the .ods writer: merged SOLL header over A–F, column
header row, position rows (Light Yellow 3 on the four money columns), a Σ row with
live SUM formulas, an optional "Nettosumme lt. Dokument" comparison row plus a red
merged warning row when the cross-sum failed, an optional Sonderpreis row, and the
percentage row. Formulas carry no cached values — Excel and LibreOffice both
recalculate an openpyxl-written .xlsx on open, which retires the .ods writer's
precompute machinery; machine readers must therefore recompute Σ from the position
rows (tests/ods_inspect.py does).

Number-format codes are locale-independent in xlsx: the same cell renders
"1.234,56 €" on a German system and "1,234.56 €" elsewhere — the entire
`number:language="de"` data-style bug class of the .ods format does not exist here.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from invoice_controller.models import (
    CrossSumCheck,
    Kostenkategorie,
    OfferDocument,
    Position,
)
from invoice_controller.narrative import render_narrative_blocks

SHEET_NAME = "Kostenaufstellung"
DESC_SHEET_NAME = "Kostenbeschreibung"

HEADER_COLS = [
    "Position",
    "Beschreibung",
    "Gesamtkosten",
    "Investitionskosten",
    "Nebenkosten",
    "Nachlass",
]

_EUR = "#,##0.00\\ €;[Red]\\-#,##0.00\\ €"
_PCT = "#,##0.00\\ %"

_YELLOW = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

_FONT = Font(name="Arial", size=10)
_FONT_BOLD = Font(name="Arial", size=10, bold=True)
_FONT_SUM = Font(name="Arial", size=10, bold=True, underline="single")
_FONT_WARN = Font(name="Arial", size=10, bold=True, color="CC0000")

_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Narrative blocks span merged rows; the estimate mirrors the .ods writer.
NAR_MIN_ROWS = 3
NAR_CHARS_PER_ROW = 50

_COL_WIDTHS = {1: 10, 2: 52, 3: 16, 4: 16, 5: 16, 6: 16}


def write_kostenaufstellung(offers: list[OfferDocument], output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    row = 1
    for idx, offer in enumerate(offers):
        if idx > 0:
            row += 1
        row = _append_offer_block(ws, offer, row)

    for col, width in _COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    _append_description_sheet(wb, offers)
    wb.save(output_path)


def _cell(ws: Worksheet, row: int, col: int, value, *, font=_FONT, align=_CENTER,
          fill=None, number_format=None) -> None:
    c = ws.cell(row, col, value)
    c.font = font
    c.alignment = align
    c.border = _BORDER
    if fill is not None:
        c.fill = fill
    if number_format is not None:
        c.number_format = number_format


def _money(ws: Worksheet, row: int, col: int, value, *, font=_FONT, fill=None) -> None:
    v = float(value) if isinstance(value, Decimal) else value
    _cell(ws, row, col, v, font=font, fill=fill, number_format=_EUR)


def _append_offer_block(ws: Worksheet, offer: OfferDocument, start_row: int) -> int:
    soll_row = start_row
    _cell(ws, soll_row, 1, _format_soll_header(offer), font=_FONT_BOLD)
    for c in range(2, 7):
        _cell(ws, soll_row, c, None, font=_FONT_BOLD)
    ws.merge_cells(start_row=soll_row, start_column=1, end_row=soll_row, end_column=6)

    col_row = soll_row + 1
    for c, name in enumerate(HEADER_COLS, start=1):
        _cell(ws, col_row, c, name, font=_FONT_BOLD)

    pos_first = col_row + 1
    for i, pos in enumerate(offer.positions):
        r = pos_first + i
        _cell(ws, r, 1, pos.pos)
        _cell(ws, r, 2, _position_label(pos), align=_LEFT)
        if pos.line_total_net is None:
            # Rate-card optional (e.g. "95,00 €/m" with no quantity): no computable
            # line total — money cells stay blank so the Σ SUM range skips them.
            for c in (3, 4, 5, 6):
                _cell(ws, r, c, None, fill=_YELLOW, number_format=_EUR)
            continue
        inv, neb, nac = _category_split(pos, pos.line_total_net)
        _money(ws, r, 3, pos.line_total_net, fill=_YELLOW)
        _money(ws, r, 4, inv, fill=_YELLOW)
        _money(ws, r, 5, neb, fill=_YELLOW)
        _money(ws, r, 6, nac, fill=_YELLOW)

    n = len(offer.positions)
    sum_row = pos_first + n
    last_pos_label = offer.positions[-1].pos if offer.positions else ""
    _cell(ws, sum_row, 1, f"Σ 1…{last_pos_label}", font=_FONT_SUM)
    _cell(ws, sum_row, 2, f"Gesamtpreis Pos.1 – Pos.{last_pos_label}", font=_FONT_SUM)
    for col_letter, col in (("C", 3), ("D", 4), ("E", 5)):
        _cell(
            ws, sum_row, col,
            f"=SUM({col_letter}{pos_first}:{col_letter}{sum_row - 1})",
            font=_FONT_SUM, number_format=_EUR,
        )
    _cell(ws, sum_row, 6, f"=C{sum_row}-D{sum_row}-E{sum_row}", font=_FONT_SUM, number_format=_EUR)

    check = offer.cross_sum
    cross_failed = not check.passed and not check.not_applicable

    next_row = sum_row + 1
    if cross_failed and check.expected is not None:
        # The document's own stated Nettosumme next to the computed Σ, so a vendor-side
        # arithmetic error (or a dropped position) is visible in the workbook itself.
        _cell(ws, next_row, 1, None, font=_FONT_BOLD)
        _cell(ws, next_row, 2, "Nettosumme lt. Dokument", font=_FONT_BOLD)
        _money(ws, next_row, 3, check.expected, font=_FONT_BOLD)
        for c in (4, 5, 6):
            _cell(ws, next_row, c, None, font=_FONT_BOLD, number_format=_EUR)
        next_row += 1

    sums = _column_sums(offer)
    sonderpreis = _effective_sonderpreis(offer)
    if sonderpreis is not None:
        spp_row = next_row
        _cell(ws, spp_row, 1, None, font=_FONT_BOLD)
        _cell(ws, spp_row, 2, f"Sonderpreis Pos.1 – Pos.{last_pos_label}", font=_FONT_BOLD)
        # C stays a literal value (not a formula): downstream readers — the F2 ratio
        # loader included — must see the negotiated price without a recalc pass.
        _money(ws, spp_row, 3, sonderpreis, font=_FONT_BOLD)
        _cell(ws, spp_row, 4, f"=D{sum_row}", font=_FONT_BOLD, number_format=_EUR)
        _cell(ws, spp_row, 5, f"=E{sum_row}", font=_FONT_BOLD, number_format=_EUR)
        _cell(ws, spp_row, 6, f"=C{sum_row}-C{spp_row}", font=_FONT_BOLD, number_format=_EUR)
        pct_row = spp_row + 1
        nachlass_ref_row = spp_row
    else:
        pct_row = next_row
        nachlass_ref_row = sum_row

    _cell(ws, pct_row, 1, None, font=_FONT_SUM)
    _cell(ws, pct_row, 2, None, font=_FONT_SUM)
    # The formulas store the plain ratio; the format's own % token scales ×100 for
    # display ("0,9184" renders as "91,84 %") — matching the consultant's sheets.
    _cell(ws, pct_row, 3, f"=C{sum_row}/C{sum_row}", number_format=_PCT)
    _cell(ws, pct_row, 4, f"=D{sum_row}/C{sum_row}", number_format=_PCT)
    _cell(ws, pct_row, 5, f"=E{sum_row}/C{sum_row}", number_format=_PCT)
    _cell(ws, pct_row, 6, f"=F{nachlass_ref_row}/C{sum_row}", number_format=_PCT)

    end_row = pct_row + 1
    if cross_failed:
        _cell(ws, end_row, 1, _cross_sum_warning_text(check), font=_FONT_WARN,
              align=_LEFT_WRAP, fill=_RED_FILL)
        for c in range(2, 7):
            _cell(ws, end_row, c, None, fill=_RED_FILL)
        ws.merge_cells(start_row=end_row, start_column=1, end_row=end_row, end_column=6)
        end_row += 1

    return end_row  # first free row after the block; caller inserts the spacer row


def _append_description_sheet(wb: openpyxl.Workbook, offers: list[OfferDocument]) -> None:
    rendered = [(offer, render_narrative_blocks(offer)) for offer in offers]
    rendered = [(offer, blocks) for offer, blocks in rendered if blocks[0] or blocks[1]]
    if not rendered:
        return

    ws = wb.create_sheet(DESC_SHEET_NAME)
    row = 1
    for idx, (offer, (inv_text, neb_text)) in enumerate(rendered):
        if idx > 0:
            row += 1
        _cell(ws, row, 1, _format_soll_header(offer), font=_FONT_BOLD)
        for c in range(2, 7):
            _cell(ws, row, c, None, font=_FONT_BOLD)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
        for text in (inv_text, neb_text):
            if not text:
                continue
            rows = max(NAR_MIN_ROWS, -(-len(text) // NAR_CHARS_PER_ROW))
            _cell(ws, row, 1, text, align=_LEFT_WRAP)
            ws.merge_cells(start_row=row, start_column=1, end_row=row + rows - 1, end_column=6)
            row += rows
    for col, width in _COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width


# ---------------------------------------------------------------------------
# Shared block semantics — same logic as the (legacy) .ods writer
# ---------------------------------------------------------------------------

def _format_soll_header(offer: OfferDocument) -> str:
    from invoice_controller.models import DocumentKind

    vendor = offer.header.vendor_short or offer.header.vendor_name
    datum = offer.header.offer_date.strftime("%d.%m.%Y")
    if offer.kind is DocumentKind.STATEMENT:
        text = f"SOLL: SCHÄTZUNG lt. Stellungnahme {vendor} vom {datum} (kein Angebot, netto angenommen)"
    else:
        text = f"SOLL: {vendor} Angebot {offer.header.offer_number} vom {datum}"
    if offer.header.title:
        text += f", {offer.header.title}"
    return text


def _position_label(pos: Position) -> str:
    if pos.optional and "option" not in pos.description.lower():
        return f"{pos.description} (optional)"
    return pos.description


def _column_sums(offer: OfferDocument) -> dict[str, Decimal]:
    gesamt = inv = neb = Decimal(0)
    for p in offer.positions:
        if p.line_total_net is None:
            continue
        gesamt += p.line_total_net
        if p.kategorie == Kostenkategorie.INVESTITIONSKOSTEN:
            inv += p.line_total_net
        elif p.kategorie == Kostenkategorie.NEBENKOSTEN:
            neb += p.line_total_net
    return {"gesamt": gesamt, "inv": inv, "neb": neb, "nac": gesamt - inv - neb}


def _effective_sonderpreis(offer: OfferDocument) -> Decimal | None:
    totals = offer.totals
    candidate: Decimal | None = None
    if totals.sonderpreis is not None:
        candidate = totals.sonderpreis
    elif totals.preisnachlass and totals.nettosumme is not None:
        candidate = totals.nettosumme - totals.preisnachlass
    if candidate is not None and candidate <= 0:
        return None
    return candidate


def _category_split(pos: Position, line_total: Decimal) -> tuple[Decimal, Decimal, Decimal]:
    zero = Decimal(0)
    if pos.kategorie == Kostenkategorie.NEBENKOSTEN:
        return zero, line_total, zero
    if pos.kategorie == Kostenkategorie.NACHLASS:
        return zero, zero, line_total
    return line_total, zero, zero


def _cross_sum_warning_text(check: CrossSumCheck) -> str:
    from invoice_controller.normalize import format_de_decimal

    if check.expected is None:
        return (
            "⚠ KREUZSUMME NICHT PRÜFBAR: keine Nettosumme im Dokument gefunden — "
            f"Σ Positionen = {format_de_decimal(check.actual)} € manuell gegen das PDF prüfen"
        )
    actual = check.actual
    if check.actual_incl_optional is not None and abs(check.actual_incl_optional - check.expected) < abs(
        actual - check.expected
    ):
        actual = check.actual_incl_optional
    diff = actual - check.expected
    return (
        f"⚠ KREUZSUMME WEICHT AB: Σ Positionen = {format_de_decimal(actual)} € vs. "
        f"Nettosumme lt. Dokument = {format_de_decimal(check.expected)} € "
        f"(Differenz {format_de_decimal(diff)} €) — Positionen manuell gegen das PDF prüfen"
    )
