"""Write the F3 Standortbeschreibung to a .docx (Word) document.

Successor of the .odt writer (format decision 2026-08-28: outputs are
Microsoft-native because most end users are on Windows). Same document shape:
company heading, address line, the section heading, the assembled prose blocks,
and the review notes.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .models import Standortbeschreibung, StandortInput

SECTION_TITLE = "1.2 Beschreibung des Standortes"


def write_standort_docx(
    result: Standortbeschreibung, inp: StandortInput, out_path: str | Path
) -> Path:
    out_path = Path(out_path)
    doc = Document()

    described = inp.betreiberfirma or inp.firma
    doc.add_heading(described, level=1)
    addr = ", ".join(p for p in (inp.strasse, f"{inp.plz} {inp.stadt}".strip()) if p.strip())
    if addr:
        doc.add_paragraph(addr)
    doc.add_heading(SECTION_TITLE, level=2)

    for block in result.text.split("\n\n"):
        for line in block.split("\n"):
            doc.add_paragraph(line)
        doc.add_paragraph("")

    if result.review_notes:
        doc.add_heading("Hinweise zur Prüfung", level=2)
        for note in result.review_notes:
            doc.add_paragraph(f"• {note}")

    doc.save(str(out_path))
    return out_path
