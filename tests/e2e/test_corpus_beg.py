"""Live BEG regression against the ground-truth Kostenzusammenstellungen.

Run: `uv run pytest tests/e2e/test_corpus_beg.py --run-live` (LLM + vision calls —
the heaviest live suite in the project; ~30 calls per EM project).

Contract (PLAN.md §Phase 2b): mechanical columns must reconcile with the
consultant's sheet; judgment cells (förderfähig) and payment cells whose proofs are
missing from the corpus are asserted FLAGGED, not equal. Row shape may differ where
the consultant folded/unfolded Schlussrechnungen — per-vendor Σ Re-Betrag is the
folding-invariant comparison.

Validated live 2026-08-31: Aepker matches row-for-row to the cent (incl. the
ZUGFeRD-parsed Zacharias invoice); Madroch reconciles per-vendor with two known,
flagged divergences (SAT brutto missing on the scan → netto fallback; Zimmerei
cumulative netto read 156,93 € above the consultant's figure — both rows carry
red flags for review).
"""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_controller.match.vendor import _overlap_score, normalize_vendor

EXAMPLES = Path(__file__).resolve().parents[2] / "examples" / "BEG"

pytestmark = pytest.mark.live

# pid → (project dir, ground-truth sheet relative to the project dir, program hint).
# The EH corpus projects ship no EKK documents → hint "eh"; Holtkamp's ground truth
# is a .xls, converted on demand (LibreOffice headless, isolated profile).
PROJECTS = {
    "EM_Aepker": ("EM_Aepker/EM", "Kostenzusammenstellung_EM_Aepker.xlsx", None),
    "EM_Madroch": ("EM_Madroch/EM Dämm. AW + Gauben", "Kostenzusammenstellung_EM_Madroch.xlsx", None),
    "EM_Schneider": ("EM_Schneider/EM LWP", "Kostenzusammenstellung_EM_Hzg_Schneider.xlsx", None),
    "EM_Buttergasse": ("EM_Buttergasse/EM", "Rechnungen/Kostenzusammenstellung_EM_Gebr._Brasseler.xlsx", None),
    "EH_Fissenewert": ("EH_Fissenewert", "Rechnungen/Kostenzusammenstellung BnD_Fissenewert.xlsx", "eh"),
    "EH_Jahnstr": ("EH_Jahnstr.16a", "geprüfte Rechnungen/Kostenzusammenstellung BnD_Jahnstraße.xlsx", "eh"),
    "EH_Holtkamp": ("EH_Holtkamp", "Rechnungen & Nachweise/Kostenzusammenstellung BnD_Holtkamp.xls", "eh"),
}


def _ground_truth_path(pid: str) -> Path:
    rel, gt, _hint = PROJECTS[pid]
    path = EXAMPLES / rel / gt
    if path.suffix.lower() != ".xls":
        return path
    import subprocess

    out_dir = Path("tmp")
    out_dir.mkdir(exist_ok=True)
    converted = out_dir / (path.stem + ".xlsx")
    if not converted.exists():
        subprocess.run(
            ["soffice", "--headless", "-env:UserInstallation=file:///tmp/claude-lo-e2e",
             "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
            check=True, capture_output=True,
        )
    return converted


def _ground_truth_rows(path: Path) -> list[tuple[str, str, Decimal]]:
    """Grouped sheets carry the Firma only on a group's first row (carry it forward),
    and the Holtkamp EH variant orders columns differently — find Re-Betrag by header."""
    ws = load_workbook(path, data_only=True)["Kostenzusammenstellung"]
    headers = {str(c.value).strip().lower(): i for i, c in enumerate(ws[2]) if c.value}
    bcol = next((i for h, i in headers.items() if h.startswith("re-betrag")), 5)
    rows, last_firma = [], None
    for row in ws.iter_rows(min_row=3):
        firma, renr, betrag = row[1].value, row[2].value, row[bcol].value
        if firma:
            last_firma = str(firma).strip()
        if last_firma and renr and isinstance(betrag, (int, float)):
            rows.append((last_firma, str(renr).strip(), Decimal(str(round(betrag, 2)))))
    return rows


def _by_vendor(rows: list[tuple[str, str, Decimal]]) -> dict[frozenset[str], Decimal]:
    sums: dict[frozenset[str], Decimal] = defaultdict(lambda: Decimal(0))
    for firma, _renr, betrag in rows:
        sums[frozenset(normalize_vendor(firma))] += betrag
    return dict(sums)


@pytest.fixture(scope="module")
def beg_results():
    from invoice_controller.extract.beg import build_kostenzusammenstellung

    from invoice_controller.beg.funding import BegProgramType

    results = {}
    for pid, (rel, _gt, hint) in PROJECTS.items():
        project = EXAMPLES / rel
        if not project.exists():
            continue
        results[pid] = build_kostenzusammenstellung(
            project,
            output_path=Path("tmp") / f"e2e_beg_{pid}.xlsx",
            program_hint=BegProgramType(hint) if hint else None,
        )
    if not results:
        pytest.skip("BEG corpus not present")
    return results


@pytest.mark.parametrize("pid", list(PROJECTS))
def test_per_vendor_sums_reconcile_or_rows_are_flagged(pid, beg_results):
    if pid not in beg_results:
        pytest.skip(f"{pid} corpus absent")
    result = beg_results[pid]
    gt_rows = _ground_truth_rows(_ground_truth_path(pid))
    gt_sums = _by_vendor(gt_rows)

    gen_rows = [(r.firma, r.re_nr, r.re_betrag or Decimal(0)) for r in result.table.rows]
    gen_sums = _by_vendor(gen_rows)
    flags_by_vendor: dict[frozenset[str], bool] = defaultdict(bool)
    for r in result.table.rows:
        key = frozenset(normalize_vendor(r.firma))
        flags_by_vendor[key] = flags_by_vendor[key] or bool(r.flags)

    for vendor_key, gt_sum in gt_sums.items():
        # Best-overlap pairing — bare token intersection over-merges on weak tokens
        # ("GmbH", "der"), seen on the EH sheets.
        scored = [(k, _overlap_score(set(k), set(vendor_key))) for k in gen_sums]
        best = max(scored, key=lambda kv: kv[1], default=(None, 0.0))
        gen = gen_sums[best[0]] if best[0] is not None and best[1] >= 0.34 else None
        if gen is None:
            # Corpus gap, not a pipeline failure: some ground-truth sheets reference
            # invoices the example folder does not ship (Fissenewert's EnergieKonzept
            # rows). Warn, don't fail — a shipped-but-unreadable file is caught by
            # result.unreadable below.
            import warnings

            warnings.warn(f"{pid}: vendor {set(vendor_key)} not in generated table (corpus gap?)")
            continue
        flagged = any(
            f for k, f in flags_by_vendor.items()
            if _overlap_score(set(k), set(vendor_key)) >= 0.34
        )
        assert abs(gen - gt_sum) <= Decimal("0.02") or flagged, (
            f"{pid}: vendor {set(vendor_key)} Σ {gen} ≠ ground truth {gt_sum} and no review flag"
        )
    assert not result.unreadable, f"{pid}: unreadable inputs {result.unreadable}"


def test_aepker_rows_match_exactly(beg_results):
    """The stronger contract for the project where extraction is fully clean
    (validated 2026-08-31): identical (Re-Nr., Re-Betrag) sets."""
    if "EM_Aepker" not in beg_results:
        pytest.skip("Aepker corpus absent")
    result = beg_results["EM_Aepker"]
    gt = {(renr, betrag) for _f, renr, betrag in _ground_truth_rows(
        _ground_truth_path("EM_Aepker")
    )}
    gen = {(r.re_nr, r.re_betrag) for r in result.table.rows}
    assert gen == gt
