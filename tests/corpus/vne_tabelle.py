"""Parse the vne-generation ``VNE-Tabelle`` .xlsx into structured ground truth.

The consultant's VNE-Tabelle is the vne-generation target artifact. Its working sheet
``(Vorlage VNE-Maske)`` holds a per-invoice table: each invoice occupies 2-3
rows (an Investitionskosten split row, a Nebenkosten split row, and a Skonto
row). We recover, per invoice: recipient, dates, window-check flag, brutto,
netto, netto-after-Skonto, and the IK / NK split amounts — plus the Σ IK / Σ NK
totals. Those are exactly the columns the vne-generation test contract asserts.

Per-vendor model: each invoice's netto is split IK/NK by a single ratio for its
vendor (taken from that vendor's cost-estimation Kostenaufstellung percentage row). This
reader does not re-derive the ratio; it reads what the consultant recorded.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

WORKSHEET = "(Vorlage VNE-Maske)"


def _dec(v: object) -> Decimal | None:
    if v is None or isinstance(v, str):
        return None
    try:
        return Decimal(str(v))
    except Exception:
        return None


def _date(v: object) -> datetime | None:
    return v if isinstance(v, datetime) else None


@dataclass
class InvoiceRow:
    recipient: str
    window_ok: str | None          # "ja" / "Ja" / "Nein"  (Bewilligungszeitraum eingehalten)
    invoice_date: datetime | None  # Rechnungsdatum / Rechnung gestellt am
    payment_date: datetime | None  # Zahlungsdatum / Ausführung
    brutto: Decimal | None
    netto: Decimal | None
    netto_after_skonto: Decimal | None
    ik_amount: Decimal = Decimal(0)     # Σ Betrag Netto Investitionskosten within the record
    nk_amount: Decimal = Decimal(0)     # Σ Betrag Netto Nebenkosten within the record
    ek_amount: Decimal = Decimal(0)     # Σ Betrag Netto Einsparkonzept (newer template only)
    categories: tuple[str, ...] = ()    # observed IK / NK / EK tags for the record


@dataclass
class VneTabelle:
    path: Path
    invoices: list[InvoiceRow] = field(default_factory=list)
    # Canonical Σ = sum of the per-invoice split columns (reliable across layouts).
    # The 'Σ gesamt' summary row is captured separately as a soft cross-check only,
    # because in some sheets it is manually offset and misaligns with the columns.
    stated_sum_ik: Decimal | None = None
    stated_sum_nk: Decimal | None = None
    header_row: int | None = None

    @property
    def sum_ik(self) -> Decimal:
        return sum((i.ik_amount for i in self.invoices), Decimal(0))

    @property
    def sum_nk(self) -> Decimal:
        return sum((i.nk_amount for i in self.invoices), Decimal(0))

    @property
    def sum_ek(self) -> Decimal:
        return sum((i.ek_amount for i in self.invoices), Decimal(0))


# The corpus ships two template generations. Rather than hardcode columns, we map
# each logical field to a column by matching header-cell text (works for both).
_HEADER_MATCHERS: dict[str, tuple[str, ...]] = {
    "recipient": ("herstellername", "empfänger", "zahlungsempfänger"),
    "inv_date": ("rechnungsdatum", "rechnung gestellt am"),
    "pay_date": ("zahlungsdatum", "ausführung"),
    "window": ("bewilligungs-zeitraum eingehalt", "bewilligungszeitraum eingehalt"),
    "brutto": ("betrag\nbrutto", "brutto"),
    "netto": ("betrag\nnetto", "betrag netto"),
    "skonto_netto": ("skto", "skonto", "rab."),
    "cat": ("ik oder", "ik \noder", "ik  \noder"),
    "ik": ("investitions-kosten", "investitions", "investitionskosten"),
    "nk": ("nebenkosten",),
    "ek": ("einsparkonzept",),
}


def _norm(s: object) -> str:
    return str(s).strip().lower() if isinstance(s, str) else ""


def _find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 40) + 1):
        row = [_norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        joined = " | ".join(row)
        if "investitions" in joined and ("rechnungsdatum" in joined or "rechnung gestellt" in joined):
            return r
    return None


def _build_colmap(ws, header_row: int) -> dict[str, int]:
    cells = {c: _norm(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)}
    colmap: dict[str, int] = {}
    for field_name, needles in _HEADER_MATCHERS.items():
        for c, text in cells.items():
            if not text:
                continue
            # "netto" must not swallow the Investitions/Nebenkosten amount columns.
            if field_name == "netto" and ("investitions" in text or "nebenkosten" in text or "skto" in text):
                continue
            if any(n in text for n in needles):
                colmap[field_name] = c
                break
    return colmap


def _find_sum_row(ws, colmap: dict[str, int]) -> tuple[Decimal | None, Decimal | None]:
    """Locate the 'Σ gesamt' totals row and return (Σ IK, Σ NK)."""
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().startswith("Σ gesamt"):
                ik = _dec(ws.cell(r, colmap["ik"]).value) if "ik" in colmap else None
                nk = _dec(ws.cell(r, colmap["nk"]).value) if "nk" in colmap else None
                return ik, nk
    return None, None


def parse_vne_tabelle(path: str | Path) -> VneTabelle:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if WORKSHEET not in wb.sheetnames:
        raise ValueError(f"{path.name}: no {WORKSHEET!r} sheet (sheets={wb.sheetnames})")
    ws = wb[WORKSHEET]
    result = VneTabelle(path=path)
    hr = _find_header_row(ws)
    result.header_row = hr
    if hr is None:
        return result
    colmap = _build_colmap(ws, hr)

    def cell(r: int, field_name: str):
        c = colmap.get(field_name)
        return ws.cell(r, c).value if c else None

    current: InvoiceRow | None = None
    for r in range(hr + 1, ws.max_row + 1):
        recipient = cell(r, "recipient")
        inv_date = _date(cell(r, "inv_date"))
        brutto = _dec(cell(r, "brutto"))

        if isinstance(recipient, str) and (
            recipient.strip().startswith("Wurde durch")
            or recipient.strip().startswith("Ergab sich")
        ):
            break

        if isinstance(recipient, str) and inv_date is not None and brutto is not None:
            current = InvoiceRow(
                recipient=recipient.replace("\n", " ").strip(),
                window_ok=(cell(r, "window") or None),
                invoice_date=inv_date,
                payment_date=_date(cell(r, "pay_date")),
                brutto=brutto,
                netto=_dec(cell(r, "netto")),
                netto_after_skonto=_dec(cell(r, "skonto_netto")),
            )
            result.invoices.append(current)

        if current is not None:
            skn = _dec(cell(r, "skonto_netto"))
            if skn is not None and current.netto_after_skonto is None:
                current.netto_after_skonto = skn
            cat = cell(r, "cat")
            if isinstance(cat, str) and cat.strip():
                current.categories = (*current.categories, cat.strip())
            ik = _dec(cell(r, "ik"))
            nk = _dec(cell(r, "nk"))
            ek = _dec(cell(r, "ek"))
            if ik is not None:
                current.ik_amount += ik
            if nk is not None:
                current.nk_amount += nk
            if ek is not None:
                current.ek_amount += ek

    result.stated_sum_ik, result.stated_sum_nk = _find_sum_row(ws, colmap)
    return result
