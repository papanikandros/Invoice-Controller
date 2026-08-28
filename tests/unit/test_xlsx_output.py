"""Kostenaufstellung .xlsx writer (template/xlsx.py) — the Microsoft-native successor
of the .ods writer. Same semantic assertions as test_ods_output.py through the shared
OfferBlock reader; Σ cells hold live formulas WITHOUT cached values (openpyxl), so Σ
assertions recompute from the position rows via block_positions_sum."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from invoice_controller.extract.cross_sum import check_offer
from invoice_controller.models import (
    Kostenkategorie,
    OfferDocument,
    OfferHeader,
    OfferTotals,
    Position,
)
from invoice_controller.template.xlsx import write_kostenaufstellung
from invoice_controller.vne.ratios import from_kostenaufstellung_xlsx
from tests.ods_inspect import block_positions_sum, read_kostenaufstellung_xlsx


def _make_offer(tmp_path: Path) -> OfferDocument:
    header = OfferHeader(
        vendor_name="Musterlieferant A GmbH",
        offer_number="0000",
        offer_date=date(2024, 4, 8),
        title="Netzanschluss",
    )
    positions = [
        Position(
            pos="1",
            description="Unterverteilung: Projektierung, Lohnkosten, Materialkosten",
            line_total_net=Decimal("18759.80"),
            kategorie=Kostenkategorie.NEBENKOSTEN,
        ),
        Position(
            pos="2",
            description="Richtpreis Montage",
            line_total_net=Decimal("89880.20"),
            kategorie=Kostenkategorie.INVESTITIONSKOSTEN,
        ),
    ]
    totals = OfferTotals(nettosumme=Decimal("108640.00"))
    cross_sum = check_offer(positions, totals)
    return OfferDocument(
        source_path=tmp_path / "muster.pdf",
        header=header,
        positions=positions,
        totals=totals,
        cross_sum=cross_sum,
    )


def test_kostenaufstellung_xlsx_roundtrips(tmp_path: Path) -> None:
    offer = _make_offer(tmp_path)
    out = tmp_path / "out.xlsx"
    write_kostenaufstellung([offer], out)

    blocks = read_kostenaufstellung_xlsx(out)
    assert len(blocks) == 1
    block = blocks[0]
    assert "Musterlieferant A" in block.soll_header
    assert "0000" in block.soll_header
    assert len(block.rows_positions) == 2

    total_col, inv_col, neb_col = 2, 3, 4

    p1_row = block.rows_positions[0]
    assert p1_row.value_at(total_col) == Decimal("18759.8")
    assert p1_row.value_at(neb_col) == Decimal("18759.8")
    assert p1_row.value_at(inv_col) == Decimal("0")

    p2_row = block.rows_positions[1]
    assert p2_row.value_at(total_col) == Decimal("89880.2")
    assert p2_row.value_at(inv_col) == Decimal("89880.2")

    assert block.row_sum is not None
    sum_cell = block.row_sum.cells[total_col]
    assert sum_cell.formula and "SUM" in sum_cell.formula
    # No cached value in the .xlsx — the machine-side Σ is recomputed from positions.
    assert block_positions_sum(block, total_col) == Decimal("108640.00")


def test_optional_positions_render_inline_tagged_and_in_sum(tmp_path: Path) -> None:
    header = OfferHeader(
        vendor_name="Pressure Company", offer_number="R-1",
        offer_date=date(2025, 11, 28), title="Druckluftstation",
    )
    positions = [
        Position(pos="1", description="BOGE S 160-4 Schraubenkompressor",
                 line_total_net=Decimal("190500.00")),
        Position(pos="2", description="BOGE S 38-4 LF Schraubenkompressor",
                 line_total_net=Decimal("23700.00"), optional=True,
                 optional_reason="Preis in Klammern"),
        Position(pos="E1", description="PVC-Verrohrung 95,00 €/m",
                 line_total_net=None, optional=True, optional_reason="Eventualposition"),
    ]
    totals = OfferTotals(nettosumme=Decimal("190500.00"))
    offer = OfferDocument(
        source_path=tmp_path / "reflex.pdf", header=header,
        positions=positions, totals=totals, cross_sum=check_offer(positions, totals),
    )
    out = tmp_path / "out.xlsx"
    write_kostenaufstellung([offer], out)

    block = read_kostenaufstellung_xlsx(out)[0]
    total_col = 2
    assert len(block.rows_positions) == 3

    priced_opt = block.rows_positions[1]
    assert "(optional)" in priced_opt.text_at(1)
    assert priced_opt.value_at(total_col) == Decimal("23700")

    rate_opt = block.rows_positions[2]
    assert "(optional)" in rate_opt.text_at(1)
    assert "95,00 €/m" in rate_opt.text_at(1)
    assert rate_opt.value_at(total_col) is None

    # The priced optional is included in the Σ range (user's chosen layout).
    assert block_positions_sum(block, total_col) == Decimal("214200")


def test_failed_cross_sum_renders_warning_and_stated_total(tmp_path: Path) -> None:
    header = OfferHeader(
        vendor_name="Musterlieferant C GmbH", offer_number="1330412",
        offer_date=date(2022, 8, 17),
    )
    positions = [
        Position(pos="1", description="Zerkleinerer", line_total_net=Decimal("373026.25")),
        Position(pos="2", description="Verschleißschutz", line_total_net=Decimal("546.25")),
    ]
    totals = OfferTotals(nettosumme=Decimal("373600.00"), preisnachlass=Decimal("5100.00"))
    cross_sum = check_offer(positions, totals)
    assert not cross_sum.passed
    offer = OfferDocument(
        source_path=tmp_path / "muster_c.pdf", header=header,
        positions=positions, totals=totals, cross_sum=cross_sum,
    )
    out = tmp_path / "out.xlsx"
    write_kostenaufstellung([offer], out)

    block = read_kostenaufstellung_xlsx(out)[0]
    assert block.row_document_total is not None, "stated-Nettosumme comparison row missing"
    assert block.row_document_total.value_at(2) == Decimal("373600")

    assert block.row_warning is not None, "cross-sum warning row missing"
    warning = block.row_warning.text_at(0)
    assert warning.startswith("⚠")
    assert "373.572,50" in warning
    assert "373.600,00" in warning
    assert "27,50" in warning

    assert len(block.rows_positions) == 2, "warning rows must not be parsed as positions"


def test_passed_cross_sum_renders_no_warning(tmp_path: Path) -> None:
    offer = _make_offer(tmp_path)
    assert offer.cross_sum.passed
    out = tmp_path / "out_ok.xlsx"
    write_kostenaufstellung([offer], out)

    block = read_kostenaufstellung_xlsx(out)[0]
    assert block.row_warning is None
    assert block.row_document_total is None


def test_f2_ratio_loader_reads_the_xlsx(tmp_path: Path) -> None:
    """The F2 ratio source must work directly off the fresh .xlsx (no recalc pass):
    the loader recomputes the block sums from the position rows."""
    offer = _make_offer(tmp_path)
    out = tmp_path / "Kostenaufstellung.xlsx"
    write_kostenaufstellung([offer], out)

    ratios = from_kostenaufstellung_xlsx(out)
    assert len(ratios) == 1
    r = ratios[0]
    assert "Musterlieferant" in r.vendor
    assert r.gesamt == Decimal("108640.00")
    assert r.beantragt_ik == Decimal("89880.20")
    assert r.beantragt_nk == Decimal("18759.80")
    assert abs(r.anteil_ik + r.anteil_nk - 1) < Decimal("0.0001")


def test_multiple_offers_render_as_separate_blocks(tmp_path: Path) -> None:
    a = _make_offer(tmp_path)
    b = _make_offer(tmp_path)
    out = tmp_path / "two.xlsx"
    write_kostenaufstellung([a, b], out)
    assert len(read_kostenaufstellung_xlsx(out)) == 2


def test_statement_block_renders_schaetzung_header(tmp_path: Path) -> None:
    from invoice_controller.extract.cross_sum import check_statement
    from invoice_controller.models import DocumentKind

    positions = [Position(pos="1", description="Netzanschluss", line_total_net=Decimal("24017.00"))]
    offer = OfferDocument(
        source_path=tmp_path / "stellungnahme.pdf",
        header=OfferHeader(vendor_name="Kunde GmbH", offer_number="", offer_date=date(2026, 1, 7)),
        positions=positions, totals=OfferTotals(),
        cross_sum=check_statement(positions),
        kind=DocumentKind.STATEMENT, vat_basis_unstated=True,
    )
    out = tmp_path / "statement.xlsx"
    write_kostenaufstellung([offer], out)
    block = read_kostenaufstellung_xlsx(out)[0]
    assert "SCHÄTZUNG lt. Stellungnahme" in block.soll_header
    # A statement's cross-sum is not_applicable — no warning row must render.
    assert block.row_warning is None


def test_narrative_renders_on_second_sheet(tmp_path: Path) -> None:
    import openpyxl

    from invoice_controller.models import CostNarrative

    offer = _make_offer(tmp_path)
    offer.narrative = CostNarrative(
        investitionskosten_items="das Roh- und Fertigmaterialhandling sowie den Extruder",
        investitionskosten_seiten=[2, 3],
        nebenkosten_items="die Montage sowie die Inbetriebnahme",
        nebenkosten_seiten=[4],
    )
    out = tmp_path / "narr.xlsx"
    write_kostenaufstellung([offer], out)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Kostenaufstellung", "Kostenbeschreibung"]
    ws = wb["Kostenbeschreibung"]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("Fertigmaterialhandling" in t for t in texts)
    assert any("Inbetriebnahme" in t for t in texts)
    # The narrative must NOT leak onto the cost sheet.
    cost_texts = [c.value for row in wb["Kostenaufstellung"].iter_rows() for c in row if isinstance(c.value, str)]
    assert not any("Fertigmaterialhandling" in t for t in cost_texts)
