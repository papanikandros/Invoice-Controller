"""Positionsabgleich — position matching (deterministic scorer), the per-vendor
Abgleich, its offer-side readers, and the second sheet of the VNE workbook.

No LLM anywhere: the scorer is pure, groups are built from constructed documents
(with_llm=False path), the writer is checked semantically incl. the Q1 rule
(any variance ≠ 0 → red fill) and the rule that the VNE-Maske sheet is untouched."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from invoice_controller.config import ProjektConfig
from invoice_controller.match.positions import Confidence, propose_matches, score_pair
from invoice_controller.models import (
    AmountCheck,
    CrossSumCheck,
    DocumentKind,
    InvoiceDocument,
    InvoicePosition,
    InvoiceType,
    Kostenkategorie,
    OfferDocument,
    OfferHeader,
    OfferTotals,
    Position,
)
from invoice_controller.template.xlsx import write_kostenaufstellung
from invoice_controller.vne.abgleich import (
    SKIPPED_NO_POSITIONS,
    AbgleichResult,
    blocks_from_kostenaufstellung_xlsx,
    blocks_from_offer_documents,
    build_abgleich,
)
from invoice_controller.vne.compute import compute_vne
from invoice_controller.vne.xlsx import ABGLEICH_SHEET_NAME, SHEET_NAME, write_vne_tabelle

RED = "00FFCCCC"


def _pos(pos: str, desc: str, total: str, *, artikel: str | None = None,
         optional: bool = False) -> Position:
    return Position(pos=pos, description=desc, line_total_net=Decimal(total),
                    artikelnummer=artikel, optional=optional,
                    kategorie=Kostenkategorie.INVESTITIONSKOSTEN)


def _ipos(desc: str, total: str, pos: str = "") -> InvoicePosition:
    return InvoicePosition(pos=pos, description=desc, line_total_net=Decimal(total))


def _offer(vendor: str, positions: list[Position]) -> OfferDocument:
    total = sum((p.line_total_net for p in positions if not p.optional), Decimal(0))
    return OfferDocument(
        source_path=Path("angebot.pdf"),
        header=OfferHeader(vendor_name=vendor, offer_number="A-1", offer_date=date(2026, 1, 10)),
        positions=positions,
        totals=OfferTotals(nettosumme=total),
        cross_sum=CrossSumCheck(expected=total, actual=total, tolerance=Decimal("0.02"), passed=True),
        kind=DocumentKind.OFFER,
    )


def _invoice(vendor: str, number: str, positions: list[InvoicePosition],
             typ: InvoiceType = InvoiceType.RECHNUNG, day: int = 1) -> InvoiceDocument:
    netto = sum((p.line_total_net for p in positions), Decimal(0))
    return InvoiceDocument(
        source_path=Path(f"{number}.pdf"), vendor_name=vendor, invoice_number=number,
        invoice_date=date(2026, 5, day), invoice_type=typ, netto=netto,
        recipient_name="Wirox GmbH",
        amount_check=AmountCheck(passed=True), positions=positions,
    )


class TestScorer:
    OFFER = [
        _pos("1", "Schraubenverdichter Typ SV-90", "42300.00", artikel="SV-90"),
        _pos("2", "Montage und Inbetriebnahme", "3800.00"),
        _pos("1-9", "Kühlturmanlagen, Pumpen, Kaltwassertank (Gesamtpaket)", "211500.00"),
    ]

    def test_exact_price_and_wording_is_high(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Verdichter SV-90 geliefert", "42300.00")])
        assert p.offer_index == 0 and p.confidence is Confidence.HIGH

    def test_article_number_carries_weak_wording(self) -> None:
        score, signals = score_pair(self.OFFER[0], _ipos("Lieferung lt. Auftrag SV-90", "41000.00"))
        assert signals["article"] > 0
        assert score >= Decimal("0.4")

    def test_partial_amount_anzahlung_matches(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Anzahlung Montage und Inbetriebnahme", "1900.00")])
        assert p.offer_index == 1
        assert p.confidence in (Confidence.HIGH, Confidence.MEDIUM)

    def test_unrelated_position_matches_nothing(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Bauzaun Miete Baustelle", "480.00")])
        assert p.offer_index is None and p.confidence is Confidence.NONE


def _lr_offer() -> OfferDocument:
    return _offer("L&R Kältetechnik", [
        _pos("1", "Schraubenverdichter Typ SV-90", "42300.00"),
        _pos("2", "Montage und Inbetriebnahme", "3800.00"),
        _pos("3", "Fernwartung REX 100", "1500.00", optional=True),
    ])


def _lr_invoices() -> list[InvoiceDocument]:
    return [
        _invoice("L&R Kältetechnik GmbH", "R-1", [
            _ipos("Schraubenverdichter SV-90", "42300.00"),      # exact
            _ipos("Montage und Inbetriebnahme", "4000.00"),       # variance +200
            _ipos("Kranstellung Sonderleistung", "750.00"),       # extra
        ]),
    ]


class TestAbgleich:
    def _result(self, extra_invoices: list[InvoiceDocument] | None = None) -> AbgleichResult:
        return build_abgleich(
            blocks_from_offer_documents([_lr_offer()]),
            _lr_invoices() + (extra_invoices or []),
            offer_source="test", with_llm=False,
        )

    def test_matching_shapes(self) -> None:
        [group] = self._result().groups
        by_pos = {r.offer_position.pos: r for r in group.rows}
        assert by_pos["1"].variance == 0
        assert by_pos["2"].variance == Decimal("200.00")
        assert not by_pos["3"].matched                      # optional, not invoiced
        assert [e.amount for e in group.extras] == [Decimal("750.00")]
        assert (group.drifted, group.missing) == (1, 0)
        # Optional positions are no part of what the vendor owed.
        assert group.offered_total == Decimal("46100.00")
        assert group.invoiced_total == Decimal("47050.00")

    def test_other_vendor_is_never_matched_across(self) -> None:
        foreign = _invoice("Bauer Elektro GmbH", "E-7", [_ipos("Schraubenverdichter SV-90", "42300.00")])
        result = self._result([foreign])
        assert result.offerless_invoices == [foreign]
        [group] = result.groups
        assert all(m.invoice is not foreign for r in group.rows for m in r.matched)

    def test_statement_offers_are_no_offer_side(self) -> None:
        statement = _lr_offer()
        statement.kind = DocumentKind.STATEMENT
        assert blocks_from_offer_documents([statement]) == []


class TestLumpSumBilling:
    def test_one_line_over_the_offer_total_spreads_pro_rata(self) -> None:
        """EK4_333 (2026-09-24): the Schlussrechnung billed the whole order as one line
        equal to Σ offer — that is a lump sum, not one variance plus N FEHLT."""
        offer = _offer("L&R Kältetechnik", [
            _pos("1-9", "Kältemaschine, Tank, Pumpen", "135700.00"),
            _pos("10", "Verrohrung und Montage", "2400.00"),
            _pos("11", "Inbetriebnahme", "2100.00"),
            _pos("12", "Fernwartung", "1500.00", optional=True),
        ])
        invoices = [
            _invoice("L & R Kältetechnik GmbH & Co.KG", "SR-1",
                     [_ipos("Auftrag 26-4132 Kälteanlage komplett", "140200.00")]),
            _invoice("L & R Kältetechnik GmbH & Co.KG", "R-2", [_ipos("Transportkosten", "700.00")]),
        ]
        [group] = build_abgleich(blocks_from_offer_documents([offer]), invoices, with_llm=False).groups
        by_pos = {r.offer_position.pos: r for r in group.rows}
        assert all(by_pos[p].variance == 0 for p in ("1-9", "10", "11"))
        assert all(m.source == "lump" and "pauschal" in m.note for m in by_pos["10"].matched)
        assert not by_pos["12"].matched                     # optional stays untouched
        assert (group.missing, group.drifted) == (0, 0)
        assert [e.amount for e in group.extras] == [Decimal("700.00")]   # transport still extra
        assert group.invoiced_total == Decimal("140900.00")


class TestKostenaufstellungOfferSide:
    def test_positions_roundtrip_through_the_cost_estimation_xlsx(self, tmp_path: Path) -> None:
        """The zero-token offer side: what cost-estimation wrote (and the consultant
        verified) reads back as matchable positions, optional flag included."""
        path = tmp_path / "Kostenaufstellung.xlsx"
        write_kostenaufstellung([_lr_offer()], path)

        [block] = blocks_from_kostenaufstellung_xlsx(path)
        assert "L&R" in block.label
        assert [(p.pos, p.line_total_net, p.optional) for p in block.positions] == [
            ("1", Decimal("42300"), False),
            ("2", Decimal("3800"), False),
            ("3", Decimal("1500"), True),
        ]

        [group] = build_abgleich([block], _lr_invoices(), with_llm=False).groups
        assert {r.offer_position.pos: r.variance for r in group.rows if r.matched} == {
            "1": Decimal(0), "2": Decimal("200"),
        }


class TestKostenaufstellungPdfOfferSide:
    LAYOUT = (
        "                Muster GmbH, Musterstadt - Vergleich Anlagen\n\n"
        "                SOLL Musterlieferant Angebot A-77 vom 07.01.2026, Anlage\n"
        "Position        Beschreibung                       Gesamtkosten   Investitionskosten   Nebenkosten\n"
        "           Maschine, Tank, Pumpen und\n"
        "  1-9      Steuerung mit Winterentlastung,       135.700,00 €     135.700,00 €      0,00 €\n"
        "           Schaltschrank und Container\n"
        "  10       Verrohrung und Montage                  2.400,00 €           0,00 €  2.400,00 €\n"
        "Σ 1…10                       Gesamtpreis Pos. 1 – 10   138.100,00 €   135.700,00 €  2.400,00 €\n"
        "  11       Filter (optional)                       3.600,00 €       3.600,00 €      0,00 €\n"
        "Σ 1…11                       Gesamtpreis Pos. 1 – 11   141.700,00 €   139.300,00 €  2.400,00 €\n"
        "                                                 100,00 %        98,31 %       1,69 %\n"
        "                SOLL Kosteneinschätzungen von Kunden\n"
        "Position        Beschreibung                       Gesamtkosten   Investitionskosten   Nebenkosten\n"
        "   1       Tiefbau                                 8.000,00 €           0,00 €  8.000,00 €\n"
        "Σ 1-1                        Gesamtpreis Pos. 1-1\n"
        "                                                   8.000,00 €           0,00 €  8.000,00 €\n"
        "                                                 100,00 %         0,00 %     100,00 %\n"
        "                                     Kostenaufstellung\n\f"
    )

    def test_positions_with_wrapped_descriptions(self) -> None:
        """EK4_333 (2026-09-24): the colleague uploads the consultant's Kostenaufstellung
        PDF, so the ratios AND the offer positions come from it — zero tokens."""
        from invoice_controller.vne.abgleich import blocks_from_layout_text

        [block] = blocks_from_layout_text(self.LAYOUT)          # Schätzung block excluded
        assert block.label == "Musterlieferant"
        assert [(p.pos, p.line_total_net, p.optional) for p in block.positions] == [
            ("1-9", Decimal("135700.00"), False),
            ("10", Decimal("2400.00"), False),
            ("11", Decimal("3600.00"), True),
        ]
        wrapped = block.positions[0].description
        assert wrapped.startswith("Maschine, Tank") and wrapped.endswith("Schaltschrank und Container")
        assert block.positions[1].description == "Verrohrung und Montage"


class TestSecondSheet:
    def _workbook(self, tmp_path: Path, abgleich: AbgleichResult | None):
        invoices = _lr_invoices()
        result = compute_vne(invoices, [], ProjektConfig(), ignored=[], ratio_source="none")
        result.abgleich = abgleich
        out = tmp_path / "vne.xlsx"
        write_vne_tabelle(result, ProjektConfig(), out)
        return load_workbook(out)

    def _abgleich(self) -> AbgleichResult:
        return build_abgleich(blocks_from_offer_documents([_lr_offer()]), _lr_invoices(),
                              offer_source="test", with_llm=False)

    def test_q1_red_on_any_variance(self, tmp_path: Path) -> None:
        wb = self._workbook(tmp_path, self._abgleich())
        assert wb.sheetnames == [SHEET_NAME, ABGLEICH_SHEET_NAME]
        ws = wb[ABGLEICH_SHEET_NAME]
        rows = {str(r[0].value): r for r in ws.iter_rows(min_row=5) if r[0].value}
        exact, drifted = rows["1"], rows["2"]
        assert exact[4].fill.start_color.rgb != RED         # variance 0 → plain
        assert drifted[4].fill.start_color.rgb == RED       # +200 € → red (Q1)
        assert rows["EXTRA"][0].fill.start_color.rgb == RED
        assert rows["3"][8].value == "optional, nicht abgerufen"
        assert rows["Σ"][4].value == 950.0                  # 47.050 − 46.100

    def test_vne_sheet_is_identical_with_and_without_abgleich(self, tmp_path: Path) -> None:
        """User requirement 2026-09-21: the VNE-Maske must match the consultant's
        examples — the Abgleich may only ever ADD a sheet."""
        def cells(wb):
            return [(c.coordinate, c.value, c.fill.start_color.rgb, c.number_format)
                    for row in wb[SHEET_NAME].iter_rows() for c in row]

        (tmp_path / "a").mkdir()
        (tmp_path / "b").mkdir()
        plain = self._workbook(tmp_path / "a", None)
        with_sheet = self._workbook(tmp_path / "b", self._abgleich())
        assert plain.sheetnames == [SHEET_NAME]
        assert cells(plain) == cells(with_sheet)

    def test_skipped_reason_is_written_not_hidden(self, tmp_path: Path) -> None:
        wb = self._workbook(tmp_path, AbgleichResult(skipped_reason=SKIPPED_NO_POSITIONS))
        assert wb[ABGLEICH_SHEET_NAME]["A2"].value == SKIPPED_NO_POSITIONS


class TestFoldingAcrossVendorSpellings:
    def test_ar_tr_and_pre_sr_gutschrift_fold_despite_name_drift(self) -> None:
        """ZePa live finding (2026-09-03): the same vendor extracts with different
        name strings across documents — folding must use overlap, and a Gutschrift
        dated before the Schlussrechnung is part of its deducted advances."""
        from invoice_controller.beg.compute import _fold_advances

        ar = _invoice("MAFAC - E. Schwarz GmbH", "2341075",
                      [_ipos("Anzahlung", "38700.00")], InvoiceType.ANZAHLUNGSRECHNUNG, day=1)
        gs = _invoice("MAFAC. E Schwarz GmbH & Co. KG", "2341322",
                      [_ipos("Gutschrift zu Anzahlung", "-5332.77")], InvoiceType.GUTSCHRIFT, day=5)
        tr = _invoice("MAFAC. - E, Schwarz GmbH & Co. KG", "2341527",
                      [_ipos("Anzahlung gemäß Vereinbarung", "41750.00")], InvoiceType.TEILRECHNUNG, day=10)
        sr = _invoice("MAFAC. - E, Schwarz GmbH & Co. KG", "2341528",
                      [_ipos("Reinigungsmaschine komplett", "83500.00")], InvoiceType.SCHLUSSRECHNUNG, day=31)
        sr.cumulative_netto = Decimal("83500.00")

        kept, folded = _fold_advances([ar, gs, tr, sr])
        assert kept == [sr]
        assert folded[id(sr)] == 3

    def test_post_sr_gutschrift_stays(self) -> None:
        from invoice_controller.beg.compute import _fold_advances

        sr = _invoice("Maler GmbH", "S-1", [_ipos("Gesamt", "10000.00")],
                      InvoiceType.SCHLUSSRECHNUNG, day=10)
        sr.cumulative_netto = Decimal("10000.00")
        gs = _invoice("Maler GmbH", "G-9", [_ipos("Nachträgliche Gutschrift", "-500.00")],
                      InvoiceType.GUTSCHRIFT, day=20)
        kept, _folded = _fold_advances([sr, gs])
        assert gs in kept


class TestOrchestratorWiring:
    def test_kostenaufstellung_xlsx_feeds_both_ratios_and_abgleich(self, tmp_path: Path, monkeypatch) -> None:
        """The zero-token path end to end: no offer is extracted, the verified
        Kostenaufstellung.xlsx serves the ratios AND the offer positions; the
        consultant's own invoice stays out of the vendor scope check."""
        from invoice_controller.extract import vne as orchestrator
        from invoice_controller.extract.classify import Classified, DocClass

        write_kostenaufstellung([_lr_offer()], tmp_path / "Kostenaufstellung.xlsx")
        own = _invoice("EnergieKonzept Krause", "EK-1", [_ipos("Einsparkonzept", "10000.00")])
        by_name = {"R-1.pdf": _lr_invoices()[0], "EK-1.pdf": own}

        monkeypatch.setattr(orchestrator, "classify_folder", lambda _d: [
            Classified(path=tmp_path / name, doc_class=DocClass.INVOICE, reason="test") for name in by_name
        ])
        monkeypatch.setattr(orchestrator, "extract_invoice", lambda path, **_kw: by_name[path.name])
        monkeypatch.setattr(orchestrator, "extract_offer",
                            lambda *_a, **_kw: (_ for _ in ()).throw(AssertionError("no offer extraction")))

        result = orchestrator.build_vne_tabelle(tmp_path, ProjektConfig(), with_llm_match=False)

        assert result.ratio_source == "xlsx:Kostenaufstellung.xlsx"
        abgleich = result.abgleich
        assert abgleich is not None and abgleich.skipped_reason is None
        assert abgleich.offer_source == "Kostenaufstellung.xlsx"
        [group] = abgleich.groups
        assert (group.drifted, len(group.extras)) == (1, 1)
        assert abgleich.offerless_invoices == []            # own-company invoice excluded, not "offerless"
