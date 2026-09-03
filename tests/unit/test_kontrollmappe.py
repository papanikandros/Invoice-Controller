"""R6 — position matching (deterministic scorer) + Kontrollmappe assembly/writer.

No LLM anywhere: the scorer is pure, groups are built from constructed documents
(with_llm=False path), the writer is checked semantically incl. the Q1 rule
(any variance ≠ 0 → red fill)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from invoice_controller.config import ClientConfig, ProjektConfig
from invoice_controller.kontrollmappe.build import (
    KontrollmappeResult,
    VendorGroup,
    _match_group,
)
from invoice_controller.kontrollmappe.xlsx import kontrollmappe_filename, write_kontrollmappe
from invoice_controller.match.positions import Confidence, propose_matches, score_pair
from invoice_controller.models import (
    AmountCheck,
    CrossSumCheck,
    DocumentKind,
    InvoiceDocument,
    InvoicePosition,
    InvoiceType,
    Kostenkategorie,
    OfferDocument,
    OfferHeader,
    OfferTotals,
    Position,
)

RED = "00FFCCCC"


def _pos(pos: str, desc: str, total: str, *, artikel: str | None = None,
         optional: bool = False) -> Position:
    return Position(pos=pos, description=desc, line_total_net=Decimal(total),
                    artikelnummer=artikel, optional=optional,
                    kategorie=Kostenkategorie.INVESTITIONSKOSTEN)


def _ipos(desc: str, total: str, pos: str = "") -> InvoicePosition:
    return InvoicePosition(pos=pos, description=desc, line_total_net=Decimal(total))


def _offer(vendor: str, positions: list[Position]) -> OfferDocument:
    total = sum((p.line_total_net for p in positions if not p.optional), Decimal(0))
    return OfferDocument(
        source_path=Path("angebot.pdf"),
        header=OfferHeader(vendor_name=vendor, offer_number="A-1", offer_date=date(2026, 1, 10)),
        positions=positions,
        totals=OfferTotals(nettosumme=total),
        cross_sum=CrossSumCheck(expected=total, actual=total, tolerance=Decimal("0.02"), passed=True),
        kind=DocumentKind.OFFER,
    )


def _invoice(vendor: str, number: str, positions: list[InvoicePosition],
             typ: InvoiceType = InvoiceType.RECHNUNG, day: int = 1) -> InvoiceDocument:
    netto = sum((p.line_total_net for p in positions), Decimal(0))
    return InvoiceDocument(
        source_path=Path(f"{number}.pdf"), vendor_name=vendor, invoice_number=number,
        invoice_date=date(2026, 5, day), invoice_type=typ, netto=netto,
        recipient_name="Wirox GmbH",
        amount_check=AmountCheck(passed=True), positions=positions,
    )


class TestScorer:
    OFFER = [
        _pos("1", "Schraubenverdichter Typ SV-90", "42300.00", artikel="SV-90"),
        _pos("2", "Montage und Inbetriebnahme", "3800.00"),
        _pos("1-9", "Kühlturmanlagen, Pumpen, Kaltwassertank (Gesamtpaket)", "211500.00"),
    ]

    def test_exact_price_and_wording_is_high(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Verdichter SV-90 geliefert", "42300.00")])
        assert p.offer_index == 0 and p.confidence is Confidence.HIGH

    def test_article_number_carries_weak_wording(self) -> None:
        score, signals = score_pair(self.OFFER[0], _ipos("Lieferung lt. Auftrag SV-90", "41000.00"))
        assert signals["article"] > 0
        assert score >= Decimal("0.4")

    def test_partial_amount_anzahlung_matches(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Anzahlung Montage und Inbetriebnahme", "1900.00")])
        assert p.offer_index == 1
        assert p.confidence in (Confidence.HIGH, Confidence.MEDIUM)

    def test_unrelated_position_matches_nothing(self) -> None:
        [p] = propose_matches(self.OFFER, [_ipos("Bauzaun Miete Baustelle", "480.00")])
        assert p.offer_index is None and p.confidence is Confidence.NONE


class TestGroupAndWriter:
    def _result(self) -> KontrollmappeResult:
        offer = _offer("L&R Kältetechnik", [
            _pos("1", "Schraubenverdichter Typ SV-90", "42300.00"),
            _pos("2", "Montage und Inbetriebnahme", "3800.00"),
            _pos("3", "Fernwartung REX 100", "1500.00", optional=True),
        ])
        invoices = [
            _invoice("L&R Kältetechnik GmbH", "R-1", [
                _ipos("Schraubenverdichter SV-90", "42300.00"),      # exact
                _ipos("Montage und Inbetriebnahme", "4000.00"),       # variance +200
                _ipos("Kranstellung Sonderleistung", "750.00"),       # extra
            ]),
        ]
        group = VendorGroup(label="L&R Kältetechnik", offers=[offer], invoices=invoices)
        _match_group(group, with_llm=False, on_progress=lambda _m: None)
        return KontrollmappeResult(
            config=ProjektConfig(
                client=ClientConfig(name="Wirox GmbH"),
                bewilligungszeitraum_start=date(2026, 1, 1),
                bewilligungszeitraum_end=date(2026, 12, 31),
            ),
            groups=[group], offerless_invoices=[], ignored=[], unreadable=[], flags=[],
        )

    def test_matching_shapes(self) -> None:
        [group] = self._result().groups
        by_pos = {r.offer_position.pos: r for r in group.rows}
        assert by_pos["1"].variance == 0
        assert by_pos["2"].variance == Decimal("200.00")
        assert not by_pos["3"].matched                      # optional, not invoiced
        assert [e.amount for e in group.extras] == [Decimal("750.00")]

    def test_writer_q1_red_on_any_variance(self, tmp_path: Path) -> None:
        out = tmp_path / "k.xlsx"
        write_kontrollmappe(self._result(), out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Übersicht", "Angebote", "Rechnungen", "Abgleich",
                                 "Prüfungen", "Summen", "Audit-Log"]
        ws = wb["Abgleich"]
        rows = {str(r[0].value): r for r in ws.iter_rows(min_row=2) if r[0].value}
        exact, drifted = rows["1"], rows["2"]
        assert exact[4].fill.start_color.rgb != RED         # variance 0 → plain
        assert drifted[4].fill.start_color.rgb == RED       # +200 € → red (Q1)
        extra = rows["EXTRA"]
        assert extra[0].fill.start_color.rgb == RED
        pruef = wb["Prüfungen"]
        first = list(pruef.iter_rows(min_row=2))[0]
        assert first[3].value == "ja"                       # date inside window
        assert first[5].value == "ja"                       # recipient = client

    def test_filename_convention(self) -> None:
        assert kontrollmappe_filename("EK4_204", date(2026, 9, 3)) == "Kontrollmappe_EK4_204_2026-09-03.xlsx"


class TestFoldingAcrossVendorSpellings:
    def test_ar_tr_and_pre_sr_gutschrift_fold_despite_name_drift(self) -> None:
        """ZePa live finding (2026-09-03): the same vendor extracts with different
        name strings across documents — folding must use overlap, and a Gutschrift
        dated before the Schlussrechnung is part of its deducted advances."""
        from invoice_controller.beg.compute import _fold_advances

        ar = _invoice("MAFAC - E. Schwarz GmbH", "2341075",
                      [_ipos("Anzahlung", "38700.00")], InvoiceType.ANZAHLUNGSRECHNUNG, day=1)
        gs = _invoice("MAFAC. E Schwarz GmbH & Co. KG", "2341322",
                      [_ipos("Gutschrift zu Anzahlung", "-5332.77")], InvoiceType.GUTSCHRIFT, day=5)
        tr = _invoice("MAFAC. - E, Schwarz GmbH & Co. KG", "2341527",
                      [_ipos("Anzahlung gemäß Vereinbarung", "41750.00")], InvoiceType.TEILRECHNUNG, day=10)
        sr = _invoice("MAFAC. - E, Schwarz GmbH & Co. KG", "2341528",
                      [_ipos("Reinigungsmaschine komplett", "83500.00")], InvoiceType.SCHLUSSRECHNUNG, day=31)
        sr.cumulative_netto = Decimal("83500.00")

        kept, folded = _fold_advances([ar, gs, tr, sr])
        assert kept == [sr]
        assert folded[id(sr)] == 3

    def test_post_sr_gutschrift_stays(self) -> None:
        from invoice_controller.beg.compute import _fold_advances

        sr = _invoice("Maler GmbH", "S-1", [_ipos("Gesamt", "10000.00")],
                      InvoiceType.SCHLUSSRECHNUNG, day=10)
        sr.cumulative_netto = Decimal("10000.00")
        gs = _invoice("Maler GmbH", "G-9", [_ipos("Nachträgliche Gutschrift", "-500.00")],
                      InvoiceType.GUTSCHRIFT, day=20)
        kept, _folded = _fold_advances([sr, gs])
        assert gs in kept
