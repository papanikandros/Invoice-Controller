"""B5 — BegTable assembly + EH/EM Kostenzusammenstellung writers (no LLM: labels
and reconciliations are constructed directly, mirroring the Madroch/Fissenewert
ground-truth patterns)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from invoice_controller.beg.compute import build_table
from invoice_controller.beg.funding import BegProgramType, ClientBasis, FundingMeta
from invoice_controller.beg.gewerk import BegSection, GewerkLabel
from invoice_controller.beg.payments import PaymentStatus, Reconciliation
from invoice_controller.beg.xlsx import write_kostenzusammenstellung
from invoice_controller.extract.beg import _dedupe
from invoice_controller.models import AmountCheck, InvoiceDocument, InvoiceType


def _invoice(
    vendor: str,
    number: str,
    netto: str,
    *,
    brutto: str | None = None,
    typ: InvoiceType = InvoiceType.RECHNUNG,
    cumulative_netto: str | None = None,
    mwst_pct: str | None = "19",
    day: int = 1,
    method: str = "pdfplumber+llm",
) -> InvoiceDocument:
    return InvoiceDocument(
        source_path=Path(f"{number}.pdf"),
        vendor_name=vendor,
        invoice_number=number,
        invoice_date=date(2025, 10, day),
        invoice_type=typ,
        netto=Decimal(netto),
        brutto=Decimal(brutto) if brutto else None,
        mwst_pct=Decimal(mwst_pct) if mwst_pct else None,
        cumulative_netto=Decimal(cumulative_netto) if cumulative_netto else None,
        amount_check=AmountCheck(passed=True),
        extraction_method=method,
    )


def _paid(inv: InvoiceDocument) -> Reconciliation:
    return Reconciliation(
        invoice=inv,
        status=PaymentStatus.PAID,
        bezahlt=inv.brutto or inv.netto,
        paid_date=date(2025, 11, 1),
        note="am 01.11.2025 bezahlt",
    )


def _meta(**overrides) -> FundingMeta:
    values = dict(
        program_type=BegProgramType.EM,
        client_basis=ClientBasis.PRIVAT,
        vorgangsnummer="94301228",
        antrag_date=date(2025, 7, 29),
        bescheid_date=date(2025, 8, 26),
        geplante_kosten_massnahmen=Decimal("60000"),
        geplante_kosten_baubegleitung=Decimal("2500"),
        foerdersatz_pct=Decimal("20"),
        foerderfaehige_kosten_cap=Decimal("60000"),
        baubegleitung_foerdersatz_pct=Decimal("50"),
        baubegleitung_kosten_cap=Decimal("5000"),
    )
    values.update(overrides)
    return FundingMeta(**values)


def _label(gewerk: str, section: BegSection = BegSection.MASSNAHME) -> GewerkLabel:
    return GewerkLabel(gewerk=gewerk, section=section)


class TestBuildTable:
    def test_schlussrechnung_folds_advances(self) -> None:
        """Madroch Zimmerei pattern: two Abschläge + Schlussrechnung with cumulative
        total → ONE row carrying the derived cumulative brutto, advances noted."""
        t1 = _invoice("Zimmerei DWB", "R-00307", "15125.32", typ=InvoiceType.TEILRECHNUNG, day=1)
        t2 = _invoice("Zimmerei DWB", "R-00317", "15603.21", typ=InvoiceType.TEILRECHNUNG, day=10)
        final = _invoice(
            "Zimmerei DWB", "R-00368", "7922.26",
            typ=InvoiceType.SCHLUSSRECHNUNG, cumulative_netto="39764.19", day=28,
        )
        invs = [t1, t2, final]
        table = build_table(
            _meta(),
            [_paid(i) for i in invs],
            {id(i): _label("Zimmerei") for i in invs},
        )
        assert len(table.massnahmen) == 1
        row = table.massnahmen[0]
        assert row.re_nr == "R-00368"
        # 39764.19 × 1.19 = 47319.39 — the consultant's Re-Betrag, derivation stated.
        assert row.re_betrag == Decimal("47319.39")
        assert any("abgeleitet" in n for n in row.anmerkung)
        assert any("inkl. 2 Abschlags" in n for n in row.anmerkung)

    def test_netto_basis_for_business_client(self) -> None:
        inv = _invoice("Brasseler", "B-1", "63452.80", brutto="75508.83")
        table = build_table(
            _meta(client_basis=ClientBasis.UNTERNEHMEN),
            [_paid(inv)],
            {id(inv): _label("WDVS")},
        )
        assert table.massnahmen[0].re_betrag == Decimal("63452.80")

    def test_no_proof_defaults_bezahlt_to_re_betrag_with_flag(self) -> None:
        inv = _invoice("Maler GmbH", "M-1", "1000.00", brutto="1190.00")
        table = build_table(
            _meta(),
            [Reconciliation(invoice=inv, status=PaymentStatus.NO_PROOF, note="kein Zahlungsnachweis")],
            {id(inv): _label("Maler")},
        )
        row = table.massnahmen[0]
        assert row.bezahlt == Decimal("1190.00")
        assert "kein Zahlungsnachweis" in row.flags

    def test_foerderfaehig_stays_empty_pending_eligibility(self) -> None:
        inv = _invoice("Maler GmbH", "M-1", "1000.00", brutto="1190.00")
        table = build_table(_meta(), [_paid(inv)], {id(inv): _label("Maler")})
        row = table.massnahmen[0]
        assert row.foerderfaehig is None
        assert "förderfähig prüfen" in row.anmerkung

    def test_sections_and_leistungszeitraum(self) -> None:
        m = _invoice("Maler GmbH", "M-1", "1000.00", brutto="1190.00", day=3)
        bb = _invoice("EnergieKonzept Krause GmbH", "RE-25/1417", "1000.00", brutto="1190.00", day=20)
        table = build_table(
            _meta(),
            [_paid(m), _paid(bb)],
            {id(m): _label("Maler"), id(bb): _label("Energieberater", BegSection.BAUBEGLEITUNG)},
        )
        assert [r.re_nr for r in table.massnahmen] == ["M-1"]
        assert [r.re_nr for r in table.baubegleitung] == ["RE-25/1417"]
        assert table.baubegleitung[0].re_positionen == "alle"
        assert table.leistungszeitraum == (date(2025, 10, 3), date(2025, 10, 20))


class TestEmWriter:
    def _table(self):
        z = _invoice("Zimmerei DWB", "R-00368", "39764.19", brutto="47319.39", day=28)
        m = _invoice("WM Maler GmbH", "453", "16377.50", brutto="19489.23", day=29)
        bb = _invoice("EnergieKonzept Krause GmbH", "RE-25/1417", "1000.00", brutto="1190.00", day=30)
        invs = [z, m, bb]
        labels = {
            id(z): _label("Zimmerei"),
            id(m): _label("Maler WDVS"),
            id(bb): _label("Energieberater", BegSection.BAUBEGLEITUNG),
        }
        return build_table(_meta(), [_paid(i) for i in invs], labels)

    def test_layout_and_formulas(self, tmp_path: Path) -> None:
        out = tmp_path / "em.xlsx"
        write_kostenzusammenstellung(self._table(), out)
        ws = load_workbook(out)[  # formulas preserved (no data_only)
            "Kostenzusammenstellung"
        ]
        assert ws.cell(1, 1).value.startswith("Kostenzusammenstellung - Technischer")
        assert [ws.cell(2, c).value for c in range(1, 11)] == [
            "Gewerk", "Firma", "Re-Nr.", "Re-Datum", "Re.-Positionen", "Re-Betrag",
            "bezahlt", "förderfähig", "Anmerkung", "Förderung",
        ]
        cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
        sums = [k for k, v in cells.items() if isinstance(v, str) and v.startswith("=SUM(F")]
        assert len(sums) == 2, "Maßnahmen + Baubegleitung Re-Betrag sums"
        assert any(isinstance(v, str) and v.startswith("=MIN(H") and "*0.2" in v for v in cells.values()), (
            "Förderung formula =MIN(HΣ,cap)×Satz"
        )
        assert "Förderung 20% auf 60.000€" in cells.values()
        assert "Förderung 50% bis 5.000€" in cells.values()
        assert "Gesamtsumme" in cells.values()
        assert "Vorgangsnummer:" in cells.values() and "94301228" in cells.values()
        assert any(isinstance(v, str) and v.startswith("30.10.2025") for v in cells.values()) or any(
            "28.10.2025-30.10.2025" in str(v) for v in cells.values()
        ), "Leistungszeiträume from invoice dates"

    def test_missing_meta_renders_fehlt(self, tmp_path: Path) -> None:
        z = _invoice("Zimmerei DWB", "R-1", "1000.00", brutto="1190.00")
        table = build_table(
            _meta(vorgangsnummer=None, foerdersatz_pct=None),
            [_paid(z)],
            {id(z): _label("Zimmerei")},
        )
        out = tmp_path / "em.xlsx"
        write_kostenzusammenstellung(table, out)
        ws = load_workbook(out)["Kostenzusammenstellung"]
        values = [c.value for row in ws.iter_rows() for c in row]
        assert "fehlt" in values
        assert "Fördersatz fehlt" in values


class TestEhWriter:
    def test_layout(self, tmp_path: Path) -> None:
        m = _invoice("Wallhorn Bedachungen", "24119", "22312.16", brutto="26551.47")
        bb = _invoice("EnergieKonzept", "RE-22/363", "1500.00", brutto="1785.00")
        table = build_table(
            _meta(program_type=BegProgramType.EH, eh_standard="EH-55-WPB", wohneinheiten=1),
            [_paid(m), _paid(bb)],
            {id(m): _label("Dachdeckung"), id(bb): _label("Energieberater", BegSection.BAUBEGLEITUNG)},
        )
        out = tmp_path / "eh.xlsx"
        write_kostenzusammenstellung(table, out)
        ws = load_workbook(out)["Kostenzusammenstellung"]
        assert ws.cell(1, 1).value.startswith("Kostenzusammenstellung - Bestätigung")
        assert [ws.cell(2, c).value for c in range(1, 9)] == [
            "Gewerk", "Firma", "Re-Nr.", "Re-Datum", "Re-Positionen", "Re-Betrag",
            "Förderfähiger Betrag", "Info",
        ]
        values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        assert "Baubegleitung & Fachplanung" in values
        assert "Summe Maßnahmen" in values
        assert "Summe Baubegleitung & Fachplanung" in values
        assert "Gesamtsumme" in values
        assert "EH-55-WPB" in values
        assert "Standard:" in values and "Wohneinheiten:" in values


class TestDedupe:
    def test_text_layer_beats_geprueft_scan(self) -> None:
        text_copy = _invoice("Maler GmbH", "M-1", "1000.00", method="pdfplumber+llm")
        scan_copy = _invoice("Maler GmbH", "M-1", "1000.00", method="vision-llm")
        kept, dropped = _dedupe([scan_copy, text_copy])
        assert kept == [text_copy]
        assert dropped == [scan_copy]

    def test_distinct_invoices_kept(self) -> None:
        a = _invoice("Maler GmbH", "M-1", "1000.00")
        b = _invoice("Maler GmbH", "M-2", "2000.00")
        kept, dropped = _dedupe([a, b])
        assert kept == [a, b] and dropped == []

    def test_same_number_and_amount_across_vendor_readings(self) -> None:
        """Madroch live finding: the same stamped invoice extracted twice with two
        different letterhead readings — caught by the (Re-Nr., netto) pass; the copy
        with the passing cross-sum wins."""
        from invoice_controller.models import CrossSumCheck

        good = _invoice("WM Maler GmbH", "453", "31655.00")
        good.position_check = CrossSumCheck(
            expected=Decimal("31655.00"), actual=Decimal("31655.00"),
            tolerance=Decimal("0.02"), passed=True,
        )
        bad = _invoice("Rlumbr & Kronshage GbR", "453", "31655.00", method="vision-llm")
        bad.position_check = CrossSumCheck(
            expected=Decimal("31655.00"), actual=Decimal("1100.00"),
            tolerance=Decimal("0.02"), passed=False,
        )
        kept, dropped = _dedupe([bad, good])
        assert kept == [good]
        assert dropped == [bad]
